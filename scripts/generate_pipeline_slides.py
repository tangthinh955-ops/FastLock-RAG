# -*- coding: utf-8 -*-
"""
Sinh slide deck 16:9 "LiveOrder — Pipeline toàn dự án" từ code thật của repo.

Chạy:  python scripts/generate_pipeline_slides.py
Xuất:  docs/LiveOrder-pipeline-slides.pptx  (13 slide)

Nội dung bám theo: lib/aho-corasick.ts, lib/live/service.ts,
lib/rag/*, lib/chat/ai.ts, lib/chat/flow.ts, prisma/schema.prisma.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "LiveOrder-pipeline-slides.pptx"

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x0F, 0x17, 0x2A)
INK = RGBColor(0x1E, 0x29, 0x3B)
SLATE = RGBColor(0x64, 0x74, 0x8B)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CANVAS = RGBColor(0xF8, 0xFA, 0xFC)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_FILL = RGBColor(0xCC, 0xFB, 0xF1)
TEAL_SOFT = RGBColor(0xE6, 0xFF, 0xFA)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
BLUE_FILL = RGBColor(0xDB, 0xEA, 0xFE)
BLUE_SOFT = RGBColor(0xEF, 0xF6, 0xFF)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_FILL = RGBColor(0xFE, 0xF3, 0xC7)
ROSE = RGBColor(0xBE, 0x12, 0x3C)
ROSE_FILL = RGBColor(0xFF, 0xE4, 0xE6)
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
PURPLE_FILL = RGBColor(0xED, 0xE9, 0xFE)
GREEN = RGBColor(0x15, 0x80, 0x3D)
GREEN_FILL = RGBColor(0xDC, 0xFC, 0xE7)

FONT = "Segoe UI"
MONO = "Consolas"
W, H = 13.333, 7.5
TOTAL = 13


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.lstrip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def _style_run(run, size, color, bold=False, italic=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def text(slide, x, y, w, h, content, size=12, color=INK, bold=False,
         font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.0, space_after=0, margins=(0.04, 0.02, 0.04, 0.02),
         italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    tf.vertical_anchor = anchor
    tf.clear()
    lines = str(content).split("\n")
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        _style_run(run, size, color, bold, italic, font)
    return box


def rich_text(slide, x, y, w, h, paragraphs, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, margins=(0.04, 0.02, 0.04, 0.02),
              line_spacing=1.0, space_after=0):
    """paragraphs: list of (text, size, color, bold, font)"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    tf.vertical_anchor = anchor
    tf.clear()
    for index, item in enumerate(paragraphs):
        value, size, color, bold, font = item
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = value
        _style_run(run, size, color, bold, False, font)
    return box


def card(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.12, line_w=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def rect(slide, x, y, w, h, fill, line=None, line_w=0.75):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def line(slide, x1, y1, x2, y2, color=SLATE, width=1.25, dash=None, arrow_head=False):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    if dash is not None:
        connector.line.dash_style = dash
    if arrow_head:
        ln = connector.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return connector


def chevron(slide, x, y, w=0.14, h=0.22, color=SLATE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def circle_number(slide, x, y, size, number, fill, text_color=WHITE, font_size=9):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(number)
    _style_run(run, font_size, text_color, True)
    return shape


def pill(slide, x, y, w, h, label, fill, text_color=WHITE, size=9, bold=True):
    card(slide, x, y, w, h, fill=fill, line=None, radius=0.5)
    text(slide, x, y, w, h, label, size=size, color=text_color, bold=bold,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def base_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, CANVAS)
    rect(slide, 0, 0, W, 0.06, TEAL)
    return slide


def header(slide, kicker, title, number):
    text(slide, 0.65, 0.30, 7.0, 0.25, kicker, size=9.5, color=TEAL, bold=True)
    text(slide, 0.63, 0.54, 11.4, 0.55, title, size=23, color=NAVY, bold=True)
    text(slide, 11.7, 0.31, 1.05, 0.25, f"{number:02d} / {TOTAL:02d}", size=9, color=MUTED,
         bold=True, align=PP_ALIGN.RIGHT)


def footer(slide, note="LiveOrder · Pipeline toàn dự án · 2026"):
    line(slide, 0.65, 7.10, 12.68, 7.10, BORDER, 0.75)
    text(slide, 0.65, 7.15, 8.0, 0.22, note, size=8, color=MUTED)
    text(slide, 9.0, 7.15, 3.68, 0.22, "Aho-Corasick  ·  RAG", size=8, color=MUTED,
         align=PP_ALIGN.RIGHT)


def step_box(slide, x, y, w, h, number, title, detail, accent, fill=WHITE,
             title_size=11.5, detail_size=8.8, num_size=9, outcome=None):
    card(slide, x, y, w, h, fill=fill, line=accent, line_w=1.1)
    circle_number(slide, x + 0.13, y + 0.13, 0.30, number, accent, font_size=num_size)
    text(slide, x + 0.50, y + 0.10, w - 0.58, 0.38, title, size=title_size,
         color=NAVY, bold=True, line_spacing=0.92)
    text(slide, x + 0.13, y + 0.55, w - 0.26, h - 0.60, detail, size=detail_size,
         color=SLATE, line_spacing=1.02)
    if outcome is not None:
        line(slide, x + 0.14, y + h - 0.47, x + w - 0.14, y + h - 0.47, accent, 0.75)
        text(slide, x + 0.13, y + h - 0.40, w - 0.26, 0.28, outcome, size=7.4,
             color=accent, bold=True)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # ============================================================ 1. TITLE
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, NAVY)
    rect(slide, 0, 0, 0.14, H, TEAL)
    rect(slide, 8.55, 0, 0.02, H, rgb("1E293B"))

    text(slide, 0.85, 1.35, 7.3, 0.3, "ĐỒ ÁN TỐT NGHIỆP · LIVESTREAM COMMERCE", size=10,
         color=rgb("5EEAD4"), bold=True)
    text(slide, 0.83, 1.85, 7.4, 1.0, "LiveOrder", size=54, color=WHITE, bold=True)
    text(slide, 0.85, 2.85, 7.3, 0.55, "PIPELINE TOÀN DỰ ÁN", size=27, color=rgb("99F6E4"),
         bold=True)
    text(slide, 0.85, 3.55, 7.2, 0.75,
         "Từ comment livestream đến đơn hàng, và AI tư vấn 24/7 —\nmột process, hai lõi thuật toán, một nguồn sự thật.",
         size=13, color=rgb("CBD5E1"), line_spacing=1.15)

    # two core chips
    card(slide, 0.85, 4.75, 3.45, 0.95, fill=rgb("0F2E2C"), line=TEAL, line_w=1.2)
    text(slide, 1.0, 4.83, 2.1, 0.25, "①  CỐT LÕI 1", size=8.5, color=rgb("5EEAD4"), bold=True)
    text(slide, 1.0, 5.10, 3.15, 0.45, "Aho-Corasick", size=17, color=WHITE, bold=True)
    text(slide, 2.75, 5.12, 1.45, 0.35, "bắt SKU → đơn", size=8.5, color=rgb("99F6E4"),
         align=PP_ALIGN.RIGHT)

    card(slide, 4.48, 4.75, 3.55, 0.95, fill=rgb("15203A"), line=rgb("3B82F6"), line_w=1.2)
    text(slide, 4.63, 4.83, 2.2, 0.25, "②  CỐT LÕI 2", size=8.5, color=rgb("93C5FD"), bold=True)
    text(slide, 4.63, 5.10, 2.0, 0.45, "RAG", size=17, color=WHITE, bold=True)
    text(slide, 5.65, 5.12, 2.2, 0.35, "Inbox AI 24/7", size=8.5, color=rgb("BFDBFE"),
         align=PP_ALIGN.RIGHT)

    text(slide, 0.85, 6.15, 7.3, 0.35,
         "Next.js 16 · React 19 · Socket.IO · Prisma 7 / PostgreSQL · e5-small local · Groq",
         size=9, color=rgb("64748B"))

    # right visual: mini pipelines
    card(slide, 9.05, 1.55, 3.65, 4.55, fill=rgb("111C30"), line=rgb("1E293B"))
    text(slide, 9.30, 1.75, 3.15, 0.25, "LIVE PIPELINE", size=8.5, color=rgb("5EEAD4"), bold=True)
    items = [
        ("comment", "SP01 2 cái"),
        ("Aho-Corasick", "match · qty · 0,x ms"),
        ("Order PENDING", "trừ kho nguyên tử"),
        ("QR → PAID", "SYSTEM message"),
    ]
    y = 2.10
    for label, detail in items:
        card(slide, 9.30, y, 3.15, 0.62, fill=rgb("17243C"), line=rgb("2B3B57"))
        text(slide, 9.42, y + 0.06, 2.95, 0.22, label, size=8.5, color=WHITE, bold=True)
        text(slide, 9.42, y + 0.29, 2.95, 0.22, detail, size=7.5, color=rgb("94A3B8"))
        if y < 4.3:
            line(slide, 10.87, y + 0.62, 10.87, y + 0.76, rgb("475569"), 1.0, arrow_head=True)
        y += 0.76

    text(slide, 9.30, 5.25, 3.15, 0.25, "AI PIPELINE", size=8.5, color=rgb("93C5FD"), bold=True)
    card(slide, 9.30, 5.52, 3.15, 0.42, fill=rgb("17243C"), line=rgb("2B3B57"))
    text(slide, 9.42, 5.58, 2.95, 0.28, "Inbox → RAG top-5 → Groq → trace", size=8,
         color=rgb("DBEAFE"))
    notes(slide,
          "Mở đầu: LiveOrder là bản demo livestream bán hàng chạy trên một process Next.js + socket.io. "
          "Bài trình bày hôm nay đi theo pipeline toàn dự án, nhưng trọng tâm là hai lõi: Aho-Corasick bắt SKU trong comment để chốt đơn, "
          "và RAG trong Inbox AI 24/7. Thông điệp xuyên suốt: mọi thứ realtime, nhưng DB là nguồn sự thật duy nhất.")


    # ============================================================ 2. BÀI TOÁN
    slide = base_slide(prs)
    header(slide, "01 · BÀI TOÁN", "Hai nỗi đau, hai lõi thuật toán", 2)

    problems = [
        ("Chốt đơn thủ công", "Comment trôi rất nhanh; bỏ sót mã, tính sai số lượng,\ntrừ kho không kịp."),
        ("Hỏi ngoài giờ", "Shop offline là khách không có ai tư vấn;\nđơn PENDING dễ bị bỏ quên."),
        ("Trả lời sai dữ liệu", "Giá, kho, chính sách thay đổi liên tục;\nAI dễ bịa nếu không có nguồn."),
    ]
    for index, (title, detail) in enumerate(problems):
        x = 0.65 + index * 4.0
        card(slide, x, 1.45, 3.75, 1.15, fill=WHITE, line=BORDER)
        rect(slide, x, 1.45, 0.07, 1.15, ROSE)
        text(slide, x + 0.18, 1.54, 3.45, 0.28, title, size=12.5, color=NAVY, bold=True)
        text(slide, x + 0.18, 1.86, 3.45, 0.62, detail, size=9, color=SLATE, line_spacing=1.0)

    # arrows down
    for index in range(3):
        x = 2.5 + index * 4.0
        line(slide, x, 2.68, x, 2.92, TEAL, 1.5, arrow_head=True)

    # two core cards
    card(slide, 0.65, 3.02, 5.95, 2.05, fill=TEAL_SOFT, line=TEAL, line_w=1.4)
    pill(slide, 0.88, 3.20, 1.25, 0.28, "CỐT LÕI 1", TEAL, size=8)
    text(slide, 0.88, 3.55, 5.35, 0.4, "Aho-Corasick", size=22, color=TEAL, bold=True)
    text(slide, 0.88, 4.02, 5.45, 0.72,
         "Bắt mọi SKU trong MỘT lượt quét; chốt đơn + trừ kho nguyên tử;\nhighlight đúng vị trí; đo thời gian thật bằng performance.now().",
         size=10.5, color=INK, line_spacing=1.12)
    text(slide, 0.88, 4.73, 5.45, 0.24,
         "trie · fail link · word boundary · atomic transaction", size=8.2, color=TEAL, bold=True)

    card(slide, 6.73, 3.02, 5.95, 2.05, fill=BLUE_SOFT, line=BLUE, line_w=1.4)
    pill(slide, 6.96, 3.20, 1.25, 0.28, "CỐT LÕI 2", BLUE, size=8)
    text(slide, 6.96, 3.55, 5.35, 0.4, "RAG", size=22, color=BLUE, bold=True)
    text(slide, 6.96, 4.02, 5.45, 0.72,
         "Index KB + sản phẩm theo từng shop; embedding 100% local; top-5\n+ giá/kho đọc tươi từ DB; Groq trả lời; trace minh bạch từng câu.",
         size=10.5, color=INK, line_spacing=1.12)
    text(slide, 6.96, 4.73, 5.45, 0.24,
         "e5-small local · cosine top-5 · context DB tươi · fallback FULL_KB", size=8.2, color=BLUE,
         bold=True)

    # outcomes
    outcomes = [
        ("Không bỏ sót đơn", "comment nào có mã là có phản hồi"),
        ("AI 24/7 đúng shop", "kể cả khi shop offline"),
        ("Kho không bao giờ âm", "hai khách tranh 1 món → 1 đơn"),
    ]
    for index, (title, detail) in enumerate(outcomes):
        x = 0.65 + index * 4.0
        card(slide, x, 5.35, 3.75, 0.95, fill=WHITE, line=BORDER)
        text(slide, x + 0.18, 5.44, 0.28, 0.28, "✓", size=14, color=GREEN, bold=True)
        text(slide, x + 0.48, 5.44, 3.1, 0.26, title, size=10.5, color=NAVY, bold=True)
        text(slide, x + 0.48, 5.73, 3.1, 0.24, detail, size=8.5, color=SLATE)
    notes(slide,
          "Ba nỗi đau: chốt đơn thủ công dễ bỏ sót, shop offline thì không ai tư vấn, và AI trả lời sai vì giá/kho thay đổi. "
          "Aho-Corasick xử lý bài toán comment → đơn; RAG xử lý bài toán tư vấn có nguồn. "
          "Kết quả mong đợi: không bỏ sót đơn, AI 24/7 đúng shop, và kho không bao giờ âm.")
    footer(slide)

    # ============================================================ 3. KIẾN TRÚC
    slide = base_slide(prs)
    header(slide, "02 · KIẾN TRÚC", "Một process, hai không gian realtime, một nguồn sự thật", 3)

    # actors
    card(slide, 0.65, 1.30, 3.7, 0.82, fill=WHITE, line=BLUE, line_w=1.2)
    text(slide, 0.80, 1.37, 1.7, 0.24, "BUYER", size=8.5, color=BLUE, bold=True)
    text(slide, 0.80, 1.60, 3.4, 0.48, "Trình duyệt · / · /live/[id] · /inbox · /orders", size=9.5,
         color=INK)

    card(slide, 8.98, 1.30, 3.7, 0.82, fill=WHITE, line=AMBER, line_w=1.2)
    text(slide, 9.13, 1.37, 1.7, 0.24, "SELLER", size=8.5, color=AMBER, bold=True)
    text(slide, 9.13, 1.60, 3.4, 0.48, "Studio · webcam preview · /seller/*", size=9.5, color=INK)

    card(slide, 4.62, 1.30, 4.1, 0.82, fill=PANEL, line=BORDER)
    text(slide, 4.77, 1.37, 3.8, 0.24, "JWT httpOnly · proxy.ts", size=8.5, color=SLATE, bold=True)
    text(slide, 4.77, 1.60, 3.8, 0.45, "Danh tính lấy từ token — không tin payload client", size=9,
         color=INK)

    line(slide, 2.5, 2.14, 2.5, 2.42, BLUE, 1.4, arrow_head=True)
    line(slide, 10.8, 2.14, 10.8, 2.42, AMBER, 1.4, arrow_head=True)

    # process container
    card(slide, 0.65, 2.44, 12.03, 2.58, fill=WHITE, line=NAVY, line_w=1.4)
    pill(slide, 5.00, 2.28, 3.35, 0.30, "MỘT PROCESS · tsx server.ts", NAVY, size=8)

    # app router
    card(slide, 0.90, 2.70, 3.2, 2.05, fill=BLUE_SOFT, line=BLUE_FILL)
    text(slide, 1.03, 2.78, 2.95, 0.26, "Next.js 16 · App Router", size=10, color=BLUE, bold=True)
    text(slide, 1.03, 3.08, 2.95, 1.55,
         "Buyer pages\nSeller pages\nREST API\nauth · products · kb\nlive · orders\nproxy.ts guard",
         size=8.3, color=INK, line_spacing=1.0)

    # socket
    card(slide, 4.35, 2.70, 3.4, 2.05, fill=PANEL, line=BORDER)
    text(slide, 4.48, 2.78, 3.1, 0.26, "Socket.IO · custom server", size=10, color=NAVY, bold=True)
    text(slide, 4.48, 3.08, 3.1, 1.55,
         "/live  join_live · send_comment\n       new_order · viewers\n/inbox  send_message\n       conversation_updated\n       ai_thinking\nJWT handshake → socket.data.user",
         size=8.0, color=INK, line_spacing=1.02)

    # engines
    card(slide, 8.05, 2.70, 4.4, 2.05, fill=TEAL_SOFT, line=TEAL_FILL)
    text(slide, 8.18, 2.78, 4.1, 0.26, "Hai engine lõi", size=10, color=TEAL, bold=True)
    card(slide, 8.18, 3.08, 4.1, 0.74, fill=WHITE, line=TEAL, line_w=1.0)
    text(slide, 8.30, 3.13, 3.9, 0.24, "① Aho-Corasick", size=9, color=TEAL, bold=True)
    text(slide, 8.30, 3.37, 3.9, 0.40, "automaton / phiên · 1 lượt quét · match + vị trí + ms", size=7.8,
         color=SLATE)
    card(slide, 8.18, 3.90, 4.1, 0.72, fill=WHITE, line=BLUE, line_w=1.0)
    text(slide, 8.30, 3.95, 3.9, 0.24, "② RAG + chat flow", size=9, color=BLUE, bold=True)
    text(slide, 8.30, 4.19, 3.9, 0.40, "index RAM / shop · router 3 nhánh · Groq · trace", size=7.8,
         color=SLATE)

    # bottom services
    card(slide, 0.65, 5.32, 5.2, 1.20, fill=WHITE, line=BORDER)
    text(slide, 0.82, 5.39, 4.9, 0.26, "PostgreSQL 16 · Docker :5433 · Prisma 7", size=9.5,
         color=NAVY, bold=True)
    text(slide, 0.82, 5.68, 4.9, 0.75,
         "User · Shop · Product · KbEntry · LiveSession\nConversation · DirectMessage · Order\nunique [shopId, sku] · eventId · dedupKey · paymentCode",
         size=7.8, color=SLATE, line_spacing=1.0)

    card(slide, 6.05, 5.32, 2.95, 1.20, fill=WHITE, line=GREEN, line_w=1.0)
    text(slide, 6.20, 5.39, 2.65, 0.26, "e5-small · local", size=9.5, color=GREEN, bold=True)
    text(slide, 6.20, 5.68, 2.65, 0.75, "multilingual · q8 · 384d\npassage: / query:\ncache 113 MB", size=7.8,
         color=SLATE, line_spacing=1.0)

    card(slide, 9.20, 5.32, 3.48, 1.20, fill=WHITE, line=PURPLE, line_w=1.0)
    text(slide, 9.35, 5.39, 3.2, 0.26, "Groq Chat Completions", size=9.5, color=PURPLE, bold=True)
    text(slide, 9.35, 5.68, 3.2, 0.75, "qwen/qwen3.8-27b\nGROQ_MODEL override\ntimeout 20s", size=7.8,
         color=SLATE, line_spacing=1.0)

    line(slide, 3.2, 5.02, 3.2, 5.30, SLATE, 1.2, arrow_head=True)
    line(slide, 7.5, 5.02, 7.5, 5.30, SLATE, 1.2, arrow_head=True)
    line(slide, 10.9, 5.02, 10.9, 5.30, SLATE, 1.2, arrow_head=True)

    text(slide, 0.65, 6.65, 12.03, 0.30,
         "globalThis giữ đúng một embedder + một index cho cả process, dù Next biên dịch code thành nhiều bundle.",
         size=8.5, color=SLATE, italic=True)
    notes(slide,
          "Điểm cần nói: chỉ một process, nhưng có hai không gian realtime — /live cho comment và đơn, /inbox cho chat và AI. "
          "Danh tính luôn lấy từ JWT ở handshake, không tin payload. PostgreSQL là nguồn sự thật duy nhất cho giá, kho, trạng thái đơn và mode hội thoại. "
          "Embedder và RAG index nằm trên globalThis để không bị nhân bản giữa các bundle của Next.")
    footer(slide)

    # ============================================================ 4. MASTER PIPELINE
    slide = base_slide(prs)
    header(slide, "03 · PIPELINE TỔNG", "Toàn cảnh: live commerce + AI 24/7 chạy song song", 4)

    pill(slide, 0.65, 1.24, 1.85, 0.28, "LIVE COMMERCE", NAVY, size=8)
    text(slide, 2.65, 1.26, 9.9, 0.25, "một vòng khép kín: mở live → chốt đơn → thanh toán → giao hàng → tổng kết",
         size=8.5, color=SLATE, italic=True)

    live_steps = [
        ("Seller mở live", "Studio webcam\nLiveSession = LIVE\nsocket live_started", "→ phiên LIVE"),
        ("Buyer comment", "socket /live · JWT\n≤200 ký tự\nche SĐT trước khi quét", "→ text đã che SĐT"),
        ("Aho-Corasick", "automaton theo phiên\n1 lượt quét\nmatch + vị trí + ms", "→ match + 0,x ms"),
        ("Order PENDING", "eventId chống trùng\nUPDATE stock ≥ qty\n+ CREATE Order", "→ Order PENDING"),
        ("QR → PAID", "QR local + paymentCode\ndemo pay → PAID\nSYSTEM message", "→ PAID + SYSTEM"),
        ("Giao hàng", "PAID → SHIPPED\nDashboard đọc DB\npopup tổng kết live", "→ Dashboard"),
    ]
    for index, (title, detail, outcome) in enumerate(live_steps):
        x = 0.65 + index * 2.03
        step_box(slide, x, 1.62, 1.85, 2.0, index + 1, title, detail,
                 [TEAL, TEAL, TEAL, AMBER, BLUE, GREEN][index], title_size=10.5, detail_size=8.0,
                 outcome=outcome)
        if index < len(live_steps) - 1:
            chevron(slide, x + 1.86, 2.49, 0.15, 0.24, MUTED)

    # shared DB band
    card(slide, 0.65, 3.82, 12.03, 0.56, fill=NAVY, line=None, radius=0.18)
    text(slide, 0.85, 3.87, 11.6, 0.22, "POSTGRESQL — NGUỒN SỰ THẬT DUY NHẤT", size=8,
         color=rgb("5EEAD4"), bold=True)
    text(slide, 0.85, 4.08, 11.65, 0.22,
         "Product(giá/kho) · Order(trạng thái) · Conversation(mode/version) · DirectMessage(trace) · LiveSession · KbEntry",
         size=8.2, color=rgb("CBD5E1"))

    pill(slide, 0.65, 4.53, 1.85, 0.28, "AI 24/7", BLUE, size=8)
    text(slide, 2.65, 4.55, 9.9, 0.25, "song song, không phụ thuộc shop online — nhưng vẫn đúng shop và đúng dữ liệu",
         size=8.5, color=SLATE, italic=True)

    ai_steps = [
        ("Buyer nhắn Inbox", "socket /inbox · JWT\nlưu BUYER + dedup\nphát realtime", "→ tin BUYER realtime"),
        ("Router", "mode = AI?\ncâu hỏi về đơn?\n3 nhánh rõ ràng", "→ 3 nhánh"),
        ("RAG top-5", "embed query local\ncosine đúng shop\nKB + sản phẩm", "→ top-5 + score"),
        ("Context + Groq", "KB top-5\nSP giá/kho tươi\n20 tin history", "→ prompt"),
        ("Answer + trace", "chip nguồn\ntopK · token · ms\nfallback FULL_KB", "→ answer + trace"),
        ("Takeover", "seller gõ tay\nAI dừng ngay\nversion check", "→ AI im lặng"),
    ]
    for index, (title, detail, outcome) in enumerate(ai_steps):
        x = 0.65 + index * 2.03
        step_box(slide, x, 4.90, 1.85, 1.62, index + 1, title, detail,
                 [BLUE, AMBER, BLUE, PURPLE, GREEN, ROSE][index], title_size=10.5, detail_size=8.0,
                 outcome=outcome)
        if index < len(ai_steps) - 1:
            chevron(slide, x + 1.86, 5.59, 0.15, 0.24, MUTED)

    text(slide, 0.65, 6.63, 12.03, 0.32,
         "Cùng một process: REST/API + socket /live + socket /inbox. JWT lấy từ handshake; DB là nơi duy nhất quyết định giá, kho và trạng thái.",
         size=8.6, color=SLATE, italic=True)
    notes(slide,
          "Đây là sơ đồ xương sống của toàn dự án. Lane trên là live commerce: seller mở live, buyer comment, Aho-Corasick bắt SKU, tạo đơn PENDING, "
          "thanh toán QR, giao hàng và tổng kết Dashboard. Lane dưới là AI 24/7 chạy song song trong Inbox: RAG top-5, Groq, trace và takeover. "
          "Hai lane gặp nhau ở DB: Product, Order, Conversation, DirectMessage.")
    footer(slide)

    # ============================================================ 5. AC ALGORITHM
    slide = base_slide(prs)
    header(slide, "04 · CỐT LÕI ①", "Aho-Corasick: bắt mọi SKU trong một lượt quét", 5)

    # trie
    card(slide, 0.65, 1.42, 5.95, 3.85, fill=WHITE, line=BORDER)
    text(slide, 0.85, 1.52, 5.5, 0.28, "Trie + fail link (BFS) + output link", size=11, color=NAVY,
         bold=True)
    text(slide, 0.85, 1.80, 5.5, 0.25, "Automaton dựng từ SKU của đúng shop đang live", size=8.5,
         color=SLATE, italic=True)

    def node(x, y, label, fill, border, width=0.62, height=0.42, size=9):
        card(slide, x, y, width, height, fill=fill, line=border, radius=0.35, line_w=1.1)
        text(slide, x, y, width, height, label, size=size, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    root_x, root_y = 0.95, 3.10
    s_x, p_x, zero_x = 1.95, 2.90, 3.85
    node(root_x, root_y, "root", BLUE_FILL, BLUE, 0.72)
    node(s_x, root_y, "S", PANEL, SLATE, 0.5)
    node(p_x, root_y, "P", PANEL, SLATE, 0.5)
    node(zero_x, root_y, "0", PANEL, SLATE, 0.5)

    line(slide, root_x + 0.72, root_y + 0.21, s_x, root_y + 0.21, SLATE, 1.1)
    line(slide, s_x + 0.5, root_y + 0.21, p_x, root_y + 0.21, SLATE, 1.1)
    line(slide, p_x + 0.5, root_y + 0.21, zero_x, root_y + 0.21, SLATE, 1.1)

    leaf_y = [1.95, 2.68, 3.41, 4.50]
    leaf_labels = ["SP01", "SP02", "SP03", "… SP07"]
    for y, label in zip(leaf_y, leaf_labels):
        node(4.85, y, label, TEAL_SOFT, TEAL, 1.05, 0.40, 8.6)
        line(slide, zero_x + 0.5, root_y + 0.21, 4.85, y + 0.20, TEAL, 1.0)

    # fail link
    line(slide, zero_x + 0.25, root_y + 0.42, root_x + 0.36, root_y + 0.42, ROSE, 1.0,
         dash=None, arrow_head=True)
    text(slide, 1.55, 3.66, 2.2, 0.22, "fail link (hậu tố dài nhất)", size=7.5, color=ROSE, italic=True)

    text(slide, 0.85, 4.78, 5.55, 0.38,
         "Không khớp → lùi theo fail link thay vì quay lại đầu chuỗi.\nOutput của nút được lan truyền theo fail link.",
         size=8.3, color=SLATE, line_spacing=1.0)

    # explanation
    card(slide, 6.78, 1.42, 5.90, 3.85, fill=WHITE, line=BORDER)
    text(slide, 6.98, 1.55, 5.5, 0.28, "Một lượt quét làm gì?", size=11, color=NAVY, bold=True)
    bullets = [
        ("Dựng trie", "mỗi SKU là một đường đi; ký tự chuẩn hoá hoa/thường."),
        ("Nối fail link", "BFS: hậu tố dài nhất của tiền tố hiện tại cũng là tiền tố khác."),
        ("Quét text", "đi theo cạnh; không có cạnh thì lùi fail — không quay lại đầu."),
        ("Chọn match", "mỗi vị trí bắt đầu chỉ giữ match DÀI NHẤT (SP010 thắng SP01)."),
        ("Word boundary", "ký tự trước/sau không thuộc [A-Z0-9_] → SP01 không khớp trong SP010."),
        ("Trả kết quả", "start/end UTF-16 trên chính chuỗi đã che → highlight không lệch."),
    ]
    y = 1.95
    for index, (title, detail) in enumerate(bullets, start=1):
        circle_number(slide, 7.0, y + 0.02, 0.25, index, TEAL, font_size=7.5)
        text(slide, 7.34, y, 1.35, 0.24, title, size=9.2, color=TEAL, bold=True)
        text(slide, 8.62, y, 3.85, 0.42, detail, size=8.8, color=INK, line_spacing=1.0)
        y += 0.51

    # complexity cards
    facts = [
        ("Build", "O(Σ |SKU|)", "dựng một lần / phiên live"),
        ("Quét", "O(n + matches)", "một lượt trên toàn comment"),
        ("Bộ nhớ", "1 automaton / phiên", "cache theo sessionId"),
        ("Thời gian", "đo thật", "performance.now(), không hard-code"),
    ]
    for index, (label, value, detail) in enumerate(facts):
        x = 0.65 + index * 3.0
        card(slide, x, 5.55, 2.88, 1.10, fill=WHITE, line=BORDER)
        text(slide, x + 0.15, 5.62, 2.55, 0.22, label.upper(), size=7.5, color=SLATE, bold=True)
        text(slide, x + 0.15, 5.86, 2.55, 0.28, value, size=13, color=TEAL, bold=True)
        text(slide, x + 0.15, 6.18, 2.55, 0.35, detail, size=7.8, color=SLATE)
    notes(slide,
          "Automaton gồm trie các SKU, fail link nối hậu tố dài nhất, và output lan truyền theo fail link. "
          "Nhờ fail link, quét text chỉ đi một lượt, độ phức tạp O(n + số match), không phải thử từng SKU. "
          "Word boundary đảm bảo SP01 không khớp trong SP010; start/end UTF-16 đảm bảo highlight đúng. "
          "Một automaton được cache theo phiên live và chỉ nạp SKU của đúng shop đó.")
    footer(slide)

    # ============================================================ 6. AC PIPELINE
    slide = base_slide(prs)
    header(slide, "05 · CỐT LÕI ①", "Comment → đơn: pipeline nguyên tử, chống trùng", 6)

    steps = [
        ("Nhận comment", "socket /live · JWT handshake\nchỉ phiên LIVE · ≤200 ký tự\nche SĐT trước khi parse", TEAL, "→ text hợp lệ"),
        ("Aho-Corasick", "automaton cache theo phiên\n1 lượt quét · match ĐẦU TIÊN\nvị trí UTF-16 + ms thật", TEAL, "→ match đầu tiên"),
        ("Chống trùng", "eventId = sha1(session|user|nonce|SKU)\nunique trong DB\nP2002 → trả đơn cũ", BLUE, "→ eventId"),
        ("Transaction", "UPDATE stock WHERE stock ≥ qty\ncount = 0 → hết hàng\nCREATE Order PENDING", AMBER, "→ Order PENDING"),
        ("Phát realtime", "comment + matches + ms\nnew_order → 2 cửa sổ\nlog Studio · hẹn nhắc tiền", GREEN, "→ 2 cửa sổ"),
    ]
    for index, (title, detail, accent, outcome) in enumerate(steps):
        x = 0.65 + index * 2.44
        step_box(slide, x, 1.42, 2.25, 2.18, index + 1, title, detail, accent,
                 title_size=11, detail_size=8.2, outcome=outcome)
        if index < len(steps) - 1:
            chevron(slide, x + 2.26, 2.40, 0.16, 0.24, MUTED)

    # code
    card(slide, 0.65, 3.85, 8.35, 1.50, fill=NAVY, line=None, radius=0.08)
    text(slide, 0.82, 3.92, 7.9, 0.22, "TRANSACTION — hai khách tranh 1 món vẫn đúng 1 người thắng", size=7.8,
         color=rgb("5EEAD4"), bold=True)
    code_lines = [
        ("BEGIN;", rgb("93C5FD")),
        ("UPDATE \"Product\" SET stock = stock - $qty", rgb("E2E8F0")),
        (" WHERE id = $id AND stock >= $qty;   -- count = 0 ⇒ hết hàng", rgb("94A3B8")),
        ("INSERT INTO \"Order\" (... status='PENDING', unitPrice, total, eventId ...);", rgb("E2E8F0")),
        ("COMMIT;", rgb("93C5FD")),
    ]
    y = 4.18
    for value, color in code_lines:
        text(slide, 0.82, y, 8.0, 0.20, value, size=8.6, color=color, font=MONO)
        y += 0.204

    card(slide, 9.22, 3.85, 3.46, 1.50, fill=ROSE_FILL, line=ROSE_FILL)
    rect(slide, 9.22, 3.85, 0.07, 1.50, ROSE)
    text(slide, 9.42, 3.94, 3.1, 0.25, "Vì sao phải nguyên tử?", size=9.5, color=ROSE, bold=True)
    text(slide, 9.42, 4.23, 3.12, 1.0,
         "Điều kiện stock ≥ qty nằm trong chính câu UPDATE.\nPostgreSQL row-lock rồi đánh giá lại → không thể cùng trừ.",
         size=8.3, color=INK, line_spacing=1.02)

    outcomes = [
        ("Tạo đơn", "Order PENDING · unitPrice snapshot · paymentCode", GREEN, GREEN_FILL),
        ("Hết hàng", "count = 0 · không ghi gì · phát \"SPxx đã hết hàng\"", AMBER, AMBER_FILL),
        ("Trùng nonce", "trả đơn cũ · duplicate: true · không phát new_order lần 2", BLUE, BLUE_FILL),
    ]
    for index, (title, detail, accent, fill) in enumerate(outcomes):
        x = 0.65 + index * 4.0
        card(slide, x, 5.62, 3.75, 0.95, fill=fill, line=fill)
        rect(slide, x, 5.62, 0.07, 0.95, accent)
        text(slide, x + 0.18, 5.70, 3.4, 0.25, title, size=10, color=accent, bold=True)
        text(slide, x + 0.18, 5.98, 3.45, 0.45, detail, size=8.0, color=INK, line_spacing=1.0)
    notes(slide,
          "Đi qua 5 bước. Điểm quan trọng nhất là bước 4: điều kiện stock >= qty nằm trong chính câu UPDATE, trong cùng transaction với INSERT Order. "
          "Vì vậy hai transaction song song không thể cùng trừ khi kho chỉ còn 1. Bước 3 chống trùng bằng eventId sinh từ session, user, nonce và SKU; "
          "request thua bị unique chặn sẽ trả đơn của request thắng. Ba kết quả cuối: tạo đơn, hết hàng, hoặc trùng nonce.")
    footer(slide)

    # ============================================================ 7. AC EVIDENCE
    slide = base_slide(prs)
    header(slide, "06 · CỐT LÕI ①", "Bằng chứng: đúng mã, đúng số lượng, đúng thời gian", 7)

    card(slide, 0.65, 1.42, 6.15, 3.62, fill=NAVY, line=None, radius=0.08)
    text(slide, 0.85, 1.55, 5.7, 0.25, "LIVE — comment thật", size=8, color=rgb("5EEAD4"), bold=True)
    terminal = [
        ("> lấy SP01 2 cái", rgb("E2E8F0"), True),
        ("", rgb("E2E8F0"), False),
        ("match    SP01          [4, 8)", rgb("99F6E4"), False),
        ("qty      2             (20 ký tự sau mã)", rgb("E2E8F0"), False),
        ("time     0.31 ms       (performance.now)", rgb("E2E8F0"), False),
        ("order    #12 · 300.000đ · PENDING", rgb("FDE68A"), False),
        ("stock    8 → 6", rgb("E2E8F0"), False),
        ("", rgb("E2E8F0"), False),
        ("// comment đã được che SĐT trước khi quét", rgb("94A3B8"), False),
    ]
    y = 1.92
    for value, color, bold in terminal:
        text(slide, 0.88, y, 5.7, 0.24, value, size=10, color=color, font=MONO, bold=bold)
        y += 0.29

    text(slide, 0.85, 4.63, 5.7, 0.28,
         "Vị trí [4, 8) trỏ đúng chuỗi phát cho client → highlight không lệch.",
         size=8.2, color=rgb("94A3B8"), italic=True)

    # table
    card(slide, 7.05, 1.42, 5.63, 3.62, fill=WHITE, line=BORDER)
    text(slide, 7.22, 1.52, 5.2, 0.26, "Các ca biên đã kiểm thử", size=10.5, color=NAVY, bold=True)
    rows = [
        ("SP01 2 cái", "✓ match + qty 2", GREEN),
        ("SP010", "✗ không khớp SP01", ROSE),
        ("xSP01", "✗ sai word boundary", ROSE),
        ("SP01SP02", "✗ dính chữ — không mã nào", ROSE),
        ("SP01 100 cái", "✓ qty 1 (số 3 chữ số)", GREEN),
        ("0912 345 678", "→ 0912***678", BLUE),
        ("2 khách · kho 1", "✓ đúng 1 đơn", GREEN),
    ]
    text(slide, 7.30, 1.84, 2.2, 0.20, "INPUT", size=7, color=SLATE, bold=True)
    text(slide, 9.45, 1.84, 2.2, 0.20, "KẾT QUẢ", size=7, color=SLATE, bold=True)
    y = 2.10
    for index, (value, result, color) in enumerate(rows):
        if index % 2 == 0:
            rect(slide, 7.20, y, 5.30, 0.36, PANEL)
        text(slide, 7.30, y + 0.07, 2.3, 0.24, value, size=8.8, color=INK, font=MONO)
        text(slide, 9.45, y + 0.07, 2.95, 0.24, result, size=8.5, color=color, bold=True)
        y += 0.39

    # bottom invariants
    invariants = [
        ("Đúng shop", "SP01 của shop A và shop B là hai sản phẩm khác nhau — automaton chỉ nạp SKU của shop phiên live."),
        ("Đúng kho", "Trừ kho trong cùng transaction với tạo đơn; điều kiện stock ≥ qty nằm trong UPDATE."),
        ("Đúng vị trí", "start/end là chỉ số UTF-16 trên chuỗi đã che — client cắt highlight từ chính text nhận được."),
    ]
    for index, (title, detail) in enumerate(invariants):
        x = 0.65 + index * 4.0
        card(slide, x, 5.26, 3.75, 1.30, fill=WHITE, line=BORDER)
        rect(slide, x, 5.26, 0.07, 1.30, TEAL)
        text(slide, x + 0.18, 5.34, 3.4, 0.26, title, size=10.5, color=TEAL, bold=True)
        text(slide, x + 0.18, 5.65, 3.45, 0.82, detail, size=8.2, color=SLATE, line_spacing=1.02)
    notes(slide,
          "Đọc ví dụ trên terminal: comment 'lấy SP01 2 cái' cho match tại [4,8), qty 2, thời gian 0,31 ms đo thật, đơn #12 và kho 8 → 6. "
          "Bảng bên phải là các ca biên đã kiểm thử: SP010, xSP01, SP01SP02 đều không khớp; số 3 chữ số không thành qty; SĐT bị che. "
          "Nhấn mạnh: đúng shop, đúng kho, đúng vị trí highlight.")
    footer(slide)

    # ============================================================ 8. RAG OVERVIEW
    slide = base_slide(prs)
    header(slide, "07 · CỐT LÕI ②", "RAG: hai pha — nạp một lần, truy vấn mỗi tin nhắn", 8)

    # phase 1
    card(slide, 0.65, 1.40, 5.90, 4.55, fill=WHITE, line=TEAL, line_w=1.2)
    pill(slide, 0.85, 1.26, 3.5, 0.30, "PHA 1 · NẠP TRI THỨC (LAZY)", TEAL, size=8)
    text(slide, 0.85, 1.52, 5.4, 0.25, "mỗi shop một index — chỉ dựng khi cần", size=8.2, color=SLATE,
         italic=True)

    phase1 = [
        ("Nguồn dữ liệu", "KbEntry(title, answer) + Product(name, sku)"),
        ("Ghép passage", "KB: title + answer\nSP: \"Sản phẩm {name} (mã {SKU})\""),
        ("Embed local", "multilingual-e5-small · q8 · 384d\nprefix passage: · mean pooling · normalize"),
        ("Index RAM", "Map<shopId> → entries + vectors\nvector normalize sẵn → cosine = dot"),
    ]
    y = 1.90
    for index, (title, detail) in enumerate(phase1):
        card(slide, 0.88, y, 5.42, 0.82, fill=TEAL_SOFT, line=TEAL_FILL)
        circle_number(slide, 1.00, y + 0.11, 0.28, index + 1, TEAL, font_size=8)
        text(slide, 1.38, y + 0.08, 4.7, 0.24, title, size=9.5, color=TEAL, bold=True)
        text(slide, 1.38, y + 0.34, 4.75, 0.44, detail, size=8.1, color=INK, line_spacing=1.0)
        if index < 3:
            line(slide, 3.59, y + 0.82, 3.59, y + 0.96, TEAL, 1.2, arrow_head=True)
        y += 0.96

    # phase 2
    card(slide, 6.78, 1.40, 5.90, 4.55, fill=WHITE, line=BLUE, line_w=1.2)
    pill(slide, 6.98, 1.26, 3.9, 0.30, "PHA 2 · TRUY VẤN (MỖI TIN NHẮN)", BLUE, size=8)
    text(slide, 6.98, 1.52, 5.4, 0.25, "câu hỏi đi qua router trước khi tới RAG", size=8.2, color=SLATE,
         italic=True)

    phase2 = [
        ("Câu hỏi", "buyer gửi qua socket /inbox · JWT"),
        ("Embed query", "cùng model · prefix query:"),
        ("Top-5", "dot với vector của ĐÚNG shop → sort giảm dần"),
        ("Dựng context", "KB top-5 + toàn bộ SP giá/kho tươi\n+ SKU exact + live đang mở + 20 tin"),
        ("Groq", "qwen3.8-27b · timeout 20s → answer + trace"),
    ]
    y = 1.90
    for index, (title, detail) in enumerate(phase2):
        card(slide, 7.00, y, 5.42, 0.68, fill=BLUE_SOFT, line=BLUE_FILL)
        circle_number(slide, 7.12, y + 0.10, 0.26, index + 1, BLUE, font_size=7.5)
        text(slide, 7.48, y + 0.06, 4.75, 0.22, title, size=9.2, color=BLUE, bold=True)
        text(slide, 7.48, y + 0.29, 4.8, 0.36, detail, size=8.0, color=INK, line_spacing=0.98)
        if index < 4:
            line(slide, 9.71, y + 0.68, 9.71, y + 0.81, BLUE, 1.2, arrow_head=True)
        y += 0.81

    # bottom band
    card(slide, 0.65, 6.15, 12.03, 0.62, fill=NAVY, line=None, radius=0.18)
    text(slide, 0.85, 6.21, 11.7, 0.22,
         "Giá/kho KHÔNG nằm trong vector — đọc tươi từ DB mỗi lần dựng prompt.",
         size=9, color=rgb("FDE68A"), bold=True)
    text(slide, 0.85, 6.44, 11.7, 0.22,
         "Sửa KB hoặc tên/SKU sản phẩm → invalidateShopIndex(shopId) → lần truy vấn kế tiếp dựng lại.",
         size=8.2, color=rgb("CBD5E1"))
    notes(slide,
          "RAG có hai pha. Pha nạp là lazy: chỉ dựng index khi shop được hỏi lần đầu, hoặc sau khi bị invalidate. "
          "Pha truy vấn đi qua router trước: nếu là câu hỏi về đơn thì đi template DB, nếu không mới embed câu hỏi và lấy top-5. "
          "Điểm mấu chốt: giá và kho không nằm trong vector — chúng được đọc tươi từ DB mỗi lần dựng prompt, nên đổi giá không cần đánh lại index.")
    footer(slide)

    # ============================================================ 9. RAG INGESTION
    slide = base_slide(prs)
    header(slide, "08 · CỐT LÕI ②", "Pha nạp: index theo shop, lazy, không vector DB", 9)

    ingest = [
        ("Dữ liệu shop", "SELECT KbEntry + Product\nWHERE shopId = phiên chat", TEAL, "→ KB + Product"),
        ("Kích hoạt", "retrieve lần đầu của shop\nhoặc sau invalidate", TEAL, "→ buildIndex()"),
        ("Ghép passage", "KB: title + answer\nSP: tên + mã SKU", BLUE, "→ passages"),
        ("Embed local", "e5-small q8 · 384d\nmean pooling · normalize", BLUE, "→ 384 chiều"),
        ("Lưu RAM", "globalThis Map<shopId>\nentries + vectors + meta", PURPLE, "→ Map<shopId>"),
        ("Vô hiệu", "KB/tên/SKU đổi\ninvalidateShopIndex(shopId)", AMBER, "→ rebuild lần sau"),
    ]
    for index, (title, detail, accent, outcome) in enumerate(ingest):
        x = 0.65 + index * 2.03
        step_box(slide, x, 1.42, 1.86, 2.05, index + 1, title, detail, accent,
                 title_size=9.8, detail_size=8.0, outcome=outcome)
        if index < len(ingest) - 1:
            chevron(slide, x + 1.87, 2.34, 0.14, 0.22, MUTED)

    facts = [
        ("1 model / process", "Singleton globalThis; cache .cache/transformers ~113 MB.\nLần đầu chậm, các lần sau local.", GREEN),
        ("Dot = cosine", "Vector đã normalize sẵn; topK mặc định = 5.\nKhông cần thư viện vector DB.", BLUE),
        ("Cô lập shop", "Index riêng theo shopId. SP01 shop A ≠ SP01 shop B.\nKhông lẫn tri thức.", TEAL),
        ("Không bền", "Restart → dựng lại. Đổi giá/kho KHÔNG invalidate\nvì giá/kho không nằm trong vector.", AMBER),
    ]
    for index, (title, detail, accent) in enumerate(facts):
        x = 0.65 + index * 3.0
        card(slide, x, 3.78, 2.88, 1.35, fill=WHITE, line=BORDER)
        rect(slide, x, 3.78, 0.07, 1.35, accent)
        text(slide, x + 0.18, 3.86, 2.55, 0.25, title, size=9.5, color=accent, bold=True)
        text(slide, x + 0.18, 4.16, 2.58, 0.88, detail, size=7.8, color=SLATE, line_spacing=1.02)

    # index record
    card(slide, 0.65, 5.35, 12.03, 1.18, fill=NAVY, line=None, radius=0.10)
    text(slide, 0.85, 5.43, 11.6, 0.22, "BẢN GHI INDEX TRONG RAM", size=8, color=rgb("5EEAD4"), bold=True)
    columns = [
        ("Map<shopId, ShopIndex>", rgb("E2E8F0")),
        ("meta: modelId · revision · builtAt\nkbCount · productCount", rgb("CBD5E1")),
        ("entries: id · kind(kb|product) · title · sku?\nvectors: number[][] · 384 chiều · normalized", rgb("CBD5E1")),
    ]
    x = 0.85
    widths = [3.2, 3.7, 4.7]
    for (value, color), width in zip(columns, widths):
        text(slide, x, 5.68, width - 0.15, 0.78, value, size=8.3, color=color, font=MONO,
             line_spacing=1.05)
        x += width
    notes(slide,
          "Index được dựng lười theo từng shop và giữ trong RAM trên globalThis. Một model e5-small duy nhất cho cả process, cache khoảng 113 MB. "
          "Vector đã normalize nên cosine chỉ là dot product, topK mặc định 5. Sửa KB hoặc tên/SKU sản phẩm sẽ gọi invalidateShopIndex; "
          "đổi giá/kho thì không cần vì hai trường này không được embed. Không có vector DB — đây là lựa chọn có chủ đích cho demo.")
    footer(slide)

    # ============================================================ 10. RAG QUERY
    slide = base_slide(prs)
    header(slide, "09 · CỐT LÕI ②", "Pha truy vấn: router 3 nhánh + top-5 + DB tươi", 10)

    query = [
        ("Nhận tin", "/inbox · JWT\nlưu BUYER · dedup nonce\nphát realtime", TEAL, "→ BUYER row"),
        ("Router", "mode = AI?\ncâu hỏi về đơn?\n3 nhánh rõ ràng", AMBER, "→ nhánh"),
        ("Retrieve", "embed query local\ndot · top-5 đúng shop\nKB + sản phẩm", BLUE, "→ top-5"),
        ("Context", "KB top-5\nSP giá/kho tươi\nSKU exact · live · 20 tin", PURPLE, "→ messages"),
        ("Groq", "qwen3.8-27b\ntimeout 20s\n1 system + history", PURPLE, "→ answer"),
        ("Lưu & phát", "AI + trace JSON\nsocket realtime\nchip nguồn + trace", GREEN, "→ UI trace"),
    ]
    for index, (title, detail, accent, outcome) in enumerate(query):
        x = 0.65 + index * 2.03
        step_box(slide, x, 1.42, 1.86, 1.90, index + 1, title, detail, accent,
                 title_size=9.8, detail_size=8.0, outcome=outcome)
        if index < len(query) - 1:
            chevron(slide, x + 1.87, 2.26, 0.14, 0.22, MUTED)

    branches = [
        ("NHÁNH ROUTER", "mode = HUMAN", "Seller đã tiếp quản → AI im lặng, không gọi model.", ROSE, ROSE_FILL),
        ("NHÁNH ROUTER", "Câu hỏi về đơn", "Template đọc DB theo buyerId + shopId. Không gọi LLM, không bịa trạng thái.", AMBER, AMBER_FILL),
        ("NHÁNH GROQ", "GroqError", "Thiếu key / HTTP lỗi / quá 20s → fallback toàn bộ KB của shop + error trong trace.", ROSE, ROSE_FILL),
        ("NHÁNH TAKE-OVER", "version đổi", "Seller vừa gõ tay → hủy kết quả AI, không lưu, không phát.", PURPLE, PURPLE_FILL),
    ]
    for index, (kicker, title, detail, accent, fill) in enumerate(branches):
        x = 0.65 + index * 3.0
        card(slide, x, 3.55, 2.88, 1.20, fill=fill, line=fill)
        rect(slide, x, 3.55, 0.07, 1.20, accent)
        text(slide, x + 0.17, 3.61, 2.55, 0.20, kicker, size=6.8, color=accent, bold=True)
        text(slide, x + 0.17, 3.81, 2.6, 0.24, title, size=9.5, color=accent, bold=True)
        text(slide, x + 0.17, 4.08, 2.58, 0.60, detail, size=7.5, color=INK, line_spacing=1.0)

    # prompt formula
    card(slide, 0.65, 4.98, 12.03, 1.65, fill=WHITE, line=BORDER)
    text(slide, 0.85, 5.08, 11.5, 0.28, "Prompt = 1 system message (vai + quy tắc + NGỮ CẢNH) + 20 tin history + câu hỏi",
         size=11, color=NAVY, bold=True)
    chips = [
        ("KB top-5", "đọc lại answer theo id", TEAL, TEAL_SOFT),
        ("Toàn bộ sản phẩm", "giá/kho tươi từ DB", AMBER, AMBER_FILL),
        ("SKU exact", "findExactSku nếu câu hỏi nêu mã", BLUE, BLUE_SOFT),
        ("Live đang mở", "tiêu đề phiên LIVE", PURPLE, PURPLE_FILL),
        ("20 tin gần nhất", "BUYER=user · AI/SELLER=assistant", GREEN, GREEN_FILL),
    ]
    for index, (title, detail, accent, fill) in enumerate(chips):
        x = 0.85 + index * 2.35
        card(slide, x, 5.48, 2.20, 0.88, fill=fill, line=fill)
        text(slide, x + 0.12, 5.55, 1.98, 0.24, title, size=8.5, color=accent, bold=True)
        text(slide, x + 0.12, 5.82, 2.0, 0.45, detail, size=7.3, color=INK, line_spacing=1.0)
    notes(slide,
          "Router có ba nhánh: nếu mode = HUMAN thì AI im lặng; nếu là câu hỏi về đơn thì trả template đọc DB theo đúng buyer + shop, không gọi LLM; "
          "còn lại mới đi RAG. Nhánh RAG lấy top-5, dựng context gồm KB, toàn bộ sản phẩm với giá/kho tươi, SKU exact, live đang mở và 20 tin gần nhất. "
          "Groq lỗi thì fallback toàn bộ KB. Takeover được kiểm tra lại bằng version sau khi AI trả lời.")
    footer(slide)

    # ============================================================ 11. PROMPT & GUARDRAILS
    slide = base_slide(prs)
    header(slide, "10 · CỐT LÕI ②", "Vì sao AI không bịa: prompt, trace và 4 chốt chặn", 11)

    card(slide, 0.65, 1.42, 7.05, 4.95, fill=NAVY, line=None, radius=0.08)
    text(slide, 0.85, 1.55, 6.6, 0.24, "SYSTEM PROMPT — DỰNG TỪ CODE THẬT", size=8, color=rgb("5EEAD4"),
         bold=True)
    prompt_lines = [
        ("Bạn là nhân viên tư vấn của shop \"…\" trên LiveOrder.", rgb("E2E8F0")),
        ("Quy tắc:", rgb("93C5FD")),
        ("- Chỉ dùng thông tin trong NGỮ CẢNH; không bịa giá/kho/chính sách.", rgb("CBD5E1")),
        ("- Ngoài ngữ cảnh → nói chưa có thông tin, mời để lại câu hỏi.", rgb("CBD5E1")),
        ("- Trả lời 1–3 câu, thân thiện, tiếng Việt; xưng \"mình\", gọi \"bạn\".", rgb("CBD5E1")),
        ("", rgb("E2E8F0")),
        ("NGỮ CẢNH", rgb("FDE68A")),
        ("- Live đang mở: \"…\" (nếu có)", rgb("CBD5E1")),
        ("- Sản phẩm: SP01 · Áo thun — 150.000đ, còn 8", rgb("CBD5E1")),
        ("- Kiến thức top-5: Bảng size: …", rgb("CBD5E1")),
        ("", rgb("E2E8F0")),
        ("HISTORY (20 tin)", rgb("93C5FD")),
        ("BUYER → user · AI/SELLER → assistant · SYSTEM bỏ", rgb("CBD5E1")),
        ("", rgb("E2E8F0")),
        ("USER  1m68 62kg mặc size gì?", rgb("E2E8F0")),
    ]
    y = 1.92
    for value, color in prompt_lines:
        text(slide, 0.88, y, 6.55, 0.22, value, size=9.2, color=color, font=MONO)
        y += 0.255

    text(slide, 0.85, 5.90, 6.6, 0.32,
         "Trace lưu vào DirectMessage.trace → UI \"Vì sao AI trả lời thế này?\"",
         size=8.6, color=rgb("94A3B8"), italic=True)

    guardrails = [
        ("Cô lập shop", "Mọi query filter shopId; index riêng. SKU nêu rõ tra chính xác bằng findExactSku — không thay bằng cosine.", TEAL, TEAL_SOFT),
        ("Đơn không bịa", "Câu hỏi đơn đi template DB theo buyerId + shopId; trạng thái luôn khớp thực tế, không gọi LLM.", AMBER, AMBER_FILL),
        ("Takeover thắng", "Chốt mode/version lúc dispatch; đọc lại sau khi AI trả lời. Lệch → hủy im lặng, không lưu, không phát.", ROSE, ROSE_FILL),
        ("Trace minh bạch", "model · topK + score · products · tokens · elapsedMs · fallback/error. Xem được ngay trên UI.", BLUE, BLUE_SOFT),
    ]
    y = 1.42
    for title, detail, accent, fill in guardrails:
        card(slide, 7.95, y, 4.73, 1.12, fill=fill, line=fill)
        rect(slide, 7.95, y, 0.07, 1.12, accent)
        text(slide, 8.12, y + 0.08, 4.35, 0.25, title, size=10, color=accent, bold=True)
        text(slide, 8.12, y + 0.35, 4.42, 0.70, detail, size=7.9, color=INK, line_spacing=1.02)
        y += 1.24

    card(slide, 0.65, 6.50, 12.03, 0.40, fill=PANEL, line=BORDER, radius=0.3)
    text(slide, 0.80, 6.55, 11.7, 0.28,
         "e5-small 384d  ·  topK 5  ·  cosine  ·  history 20  ·  qwen3.8-27b  ·  timeout 20s  ·  fallback FULL_KB",
         size=8.6, color=SLATE, bold=True, align=PP_ALIGN.CENTER)
    notes(slide,
          "Prompt chỉ có một system message chứa vai, quy tắc và toàn bộ ngữ cảnh; sau đó là lịch sử 20 tin và câu hỏi. "
          "Bốn chốt chặn: cô lập shop, câu hỏi đơn đi template DB, takeover thắng bằng version, và trace đầy đủ để giải thích. "
          "Trace lưu model, topK + score, sản phẩm, token, thời gian, fallback/error — hiển thị ở UI 'Vì sao AI trả lời thế này?'.")
    footer(slide)

    # ============================================================ 12. RUN TIMELINE
    slide = base_slide(prs)
    header(slide, "11 · VẬN HÀNH", "Hai lõi gặp nhau trong một phiên demo 10 phút", 12)

    timeline_y = 3.85
    line(slide, 0.90, timeline_y, 12.55, timeline_y, BORDER, 3.0)
    moments = [
        ("0:00", "Seller mở live", "Studio webcam · LiveSession = LIVE · socket live_started.", AMBER, "→ phiên LIVE"),
        ("0:30", "SP01 2 cái", "Aho-Corasick bắt mã · Order PENDING · trừ kho · toast 2 cửa sổ.", TEAL, "→ Order PENDING"),
        ("1:30", "Hỏi size", "RAG top-5 bảng size · Groq trả lời · chip nguồn + trace.", BLUE, "→ chip nguồn + trace"),
        ("2:00", "Tra đơn", "Template DB đúng buyerId + shopId · không gọi LLM.", AMBER, "→ trạng thái đúng DB"),
        ("6:00", "QR → PAID", "QR local + paymentCode · SYSTEM message · 2 inbox nhảy.", GREEN, "→ PAID + SYSTEM"),
        ("8:00", "Takeover & kết thúc", "Seller gõ tay → AI im · popup tổng kết → Dashboard.", ROSE, "→ Dashboard"),
    ]
    for index, (time, title, detail, accent, outcome) in enumerate(moments):
        x = 0.65 + index * 1.97
        above = index % 2 == 0
        card_y = 1.55 if above else 4.55
        card(slide, x, card_y, 1.92, 1.55, fill=WHITE, line=accent, line_w=1.1)
        text(slide, x + 0.12, card_y + 0.09, 1.65, 0.24, time, size=9, color=accent, bold=True)
        text(slide, x + 0.12, card_y + 0.36, 1.70, 0.30, title, size=9.5, color=NAVY, bold=True)
        text(slide, x + 0.12, card_y + 0.72, 1.70, 0.55, detail, size=7.2, color=SLATE,
             line_spacing=1.0)
        text(slide, x + 0.12, card_y + 1.25, 1.70, 0.22, outcome, size=7.2, color=accent, bold=True)
        center_x = x + 0.96
        if above:
            line(slide, center_x, card_y + 1.55, center_x, timeline_y - 0.06, accent, 1.4)
            circle_number(slide, center_x - 0.075, timeline_y - 0.075, 0.15, "", accent, font_size=6)
        else:
            line(slide, center_x, timeline_y + 0.06, center_x, card_y, accent, 1.4)
            circle_number(slide, center_x - 0.075, timeline_y - 0.075, 0.15, "", accent, font_size=6)

    text(slide, 0.65, 6.38, 12.03, 0.42,
         "Mỗi mốc đều để lại bằng chứng: log ms ở Studio, chip nguồn + trace ở Inbox, trạng thái đơn trong DB và Dashboard.",
         size=9.2, color=SLATE, italic=True, align=PP_ALIGN.CENTER)
    notes(slide,
          "Đây là kịch bản demo 10 phút. Mở live, khách chốt SP01 2 cái để thấy Aho-Corasick; hỏi size để thấy RAG; hỏi đơn để thấy template DB; "
          "thanh toán QR để thấy SYSTEM message và inbox nhảy; cuối cùng seller takeover để thấy AI dừng, rồi kết thúc live và mở Dashboard. "
          "Mỗi mốc đều có bằng chứng: log ms, chip nguồn, trace, trạng thái đơn trong DB.")
    footer(slide)

    # ============================================================ 13. SUMMARY
    slide = base_slide(prs)
    header(slide, "12 · TỔNG KẾT", "Điều gì làm nên pipeline này", 13)

    # left
    card(slide, 0.65, 1.42, 5.95, 3.15, fill=TEAL_SOFT, line=TEAL, line_w=1.2)
    text(slide, 0.88, 1.55, 5.3, 0.30, "① Aho-Corasick", size=16, color=TEAL, bold=True)
    text(slide, 0.88, 1.88, 5.3, 0.24, "comment → SKU → đơn", size=8.5, color=TEAL, bold=True)
    metrics_a = [
        ("1 lượt quét", "cho mỗi comment"),
        ("O(n + matches)", "không quét lại theo từng SKU"),
        ("0,x ms", "đo bằng performance.now()"),
        ("Kho nguyên tử", "UPDATE … WHERE stock >= qty"),
    ]
    y = 2.25
    for label, detail in metrics_a:
        text(slide, 0.90, y, 1.95, 0.24, label, size=9.5, color=NAVY, bold=True)
        text(slide, 2.85, y, 3.45, 0.24, detail, size=8.4, color=SLATE)
        line(slide, 0.90, y + 0.36, 6.25, y + 0.36, TEAL_FILL, 0.75)
        y += 0.48

    # right
    card(slide, 6.73, 1.42, 5.95, 3.15, fill=BLUE_SOFT, line=BLUE, line_w=1.2)
    text(slide, 6.96, 1.55, 5.3, 0.30, "② RAG", size=16, color=BLUE, bold=True)
    text(slide, 6.96, 1.88, 5.3, 0.24, "tri thức shop → câu trả lời", size=8.5, color=BLUE, bold=True)
    metrics_b = [
        ("top-5", "cosine trên vector đã normalize"),
        ("384 chiều", "e5-small, embedding 100% local"),
        ("20 giây", "timeout Groq + fallback toàn KB"),
        ("Trace", "model · nguồn · token · thời gian"),
    ]
    y = 2.25
    for label, detail in metrics_b:
        text(slide, 6.98, y, 1.75, 0.24, label, size=9.5, color=NAVY, bold=True)
        text(slide, 8.75, y, 3.65, 0.24, detail, size=8.4, color=SLATE)
        line(slide, 6.98, y + 0.36, 12.40, y + 0.36, BLUE_FILL, 0.75)
        y += 0.48

    # three closing statements
    statements = [
        ("Aho-Corasick", "biến comment thành đơn — nhanh, đúng, không oversell.", TEAL, TEAL_SOFT),
        ("RAG", "biến dữ liệu shop thành câu trả lời — đúng shop, đúng giá/kho, có dẫn nguồn.", BLUE, BLUE_SOFT),
        ("Niềm tin", "takeover + trace + fallback: người thật luôn thắng, AI không bịa.", PURPLE, PURPLE_FILL),
    ]
    for index, (title, detail, accent, fill) in enumerate(statements):
        x = 0.65 + index * 4.0
        card(slide, x, 4.85, 3.75, 1.35, fill=fill, line=fill)
        circle_number(slide, x + 0.16, 4.98, 0.30, index + 1, accent, font_size=9)
        text(slide, x + 0.56, 4.97, 3.0, 0.26, title, size=10.5, color=accent, bold=True)
        text(slide, x + 0.56, 5.28, 3.0, 0.78, detail, size=8.5, color=INK, line_spacing=1.05)

    text(slide, 0.65, 6.48, 12.03, 0.30,
         "Next.js 16 · React 19 · Socket.IO · Prisma 7 / PostgreSQL · Transformers.js e5-small · Groq",
         size=8.8, color=SLATE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, 0.65, 6.82, 12.03, 0.22, "Cảm ơn hội đồng — Q&A", size=11, color=TEAL, bold=True,
         align=PP_ALIGN.CENTER)
    notes(slide,
          "Chốt ba ý: Aho-Corasick biến comment thành đơn nhanh, đúng và không oversell; RAG biến dữ liệu shop thành câu trả lời đúng shop, đúng giá/kho, có dẫn nguồn; "
          "takeover + trace + fallback giữ niềm tin — người thật luôn thắng, AI không bịa. "
          "Cảm ơn hội đồng, mời đặt câu hỏi.")
    footer(slide)

    prs.core_properties.title = "LiveOrder — Pipeline toàn dự án (Aho-Corasick & RAG)"
    prs.core_properties.subject = "Đồ án tốt nghiệp — livestream commerce"
    prs.core_properties.author = "LiveOrder"
    prs.save(OUT)
    print(f"Saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
